"""
build_docx.py -- assemble the PeerJ Computer Science submission from paper/draft/*.md.
Writes no prose.

Outputs
  paper/manuscript.docx                 the manuscript, text only
  paper/submission/Figure N.png / .pdf  figures, numbered by first mention (PNG >= 3000 px)
  paper/submission/figure_legends.docx  one title per figure
  paper/submission/Table N.docx         one table per file, title and units in the headers
  paper/submission/cover_page.docx      title, authors, affiliation, corresponding author
  paper/submission/cover_letter.docx    the cover letter, from paper/draft/cover_letter.md
  paper/submission/manuscript.docx      a copy of paper/manuscript.docx
  paper/submission/<SUPPLEMENTARY>      supplementary data files named in table footnotes

Sections, in order: Abstract (run-in headings Background, Methods, Results, Conclusions);
1 Background; 2 Related Work; 3 Methods; 4 Results; 5 Discussion; 6 Threats to Validity;
7 Audit and Deviations from Preregistration; 8 Conclusions; Funding; Competing Interests;
Data Availability; AI-Use Statement; References. Markdown "- " items render as bullets. A section with a prose draft
is rendered from it. A section with only an outline file, or no file, gets its heading and
the visible placeholder [TO WRITE]; nothing is generated to fill it.

Rendering rules
  * Every %[...] drafting marker and every "> **Drafting note" blockquote is removed.
  * Figures and tables are not embedded. They are numbered in order of first mention and
    every in-text mention is renumbered to match: "Table N", and for figures "Fig. N"
    ("Figs. N and M") in running text, "Figure N" at the start of a sentence. Draft numbers
    resolve to files through the reference map in the results draft's drafting note;
    titles come from paper/tables/tables_index.csv and paper/figures/figures_index.csv.
  * Citations are written [@key] or [@key1; @key2]. They render name-year, e.g.
    (Smith, 2010), (Smith & Jones, 2010), (Smith, Jones & Brown, 2010),
    (Smith et al., 2010), several separated by semicolons in chronological order. A
    citation renders only if every key in it is a row of paper/references.csv with status
    VERIFIED; otherwise the bracket, like any [CITATION NEEDED: ...], stays as visible
    text. The reference list holds every VERIFIED row, alphabetically by first author,
    with its DOI, or its URL where it has no DOI. The build reports VERIFIED rows the
    rendered text does not cite.
  * Format: US Letter, 2.5 cm margins, Times New Roman 12 pt, 1.5 line spacing, left
    aligned, continuous line numbering, Word's built-in Title/Heading styles, Normal style
    for body text.
  * The build fails if the abstract exceeds 500 words.

Run: python paper/build_docx.py
"""
from __future__ import annotations

import csv
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
DRAFT = os.path.join(HERE, "draft")
OUT = os.path.join(HERE, "manuscript.docx")
SUB = os.path.join(HERE, "submission")
REFS = os.path.join(HERE, "references.csv")
FONT = "Times New Roman"
ABSTRACT_MAX_WORDS = 500
MIN_FIGURE_PX = 3000

ABSTRACT_FILE = "abstract.md"          # prose abstract, when it exists
COVER_LETTER = "cover_letter.md"       # rendered to submission/, not into the manuscript
# Written into submission/ by a later step (paper/make_checklist.py), not stale.
LATER_FILES = {"SUBMISSION_CHECKLIST.md", "FORM_ANSWERS.md"}
ABSTRACT_PARTS = ["Background.", "Methods.", "Results.", "Conclusions."]

# (heading used when there is no prose, prose draft, outline draft)
SECTIONS = [
    ("Abstract", ABSTRACT_FILE, "outline_abstract.md"),
    ("1. Background", "background.md", "outline_intro.md"),
    ("2. Related Work", "related_work.md", "outline_related_work.md"),
    ("3. Methods", "method.md", None),
    ("4. Results", "results.md", None),
    ("5. Discussion", "discussion.md", "outline_discussion.md"),
    ("6. Threats to Validity", "limitations.md", None),
    ("7. Audit and Deviations from Preregistration", "audit_disclosure.md", None),
    ("8. Conclusions", "conclusions.md", "outline_conclusion.md"),
    ("Funding", "funding.md", None),
    ("Competing Interests", "competing_interests.md", None),
    ("Data Availability", "data_availability.md", None),
    ("AI-Use Statement", "ai_statement.md", None),
]
TRAILING: list = []                     # sections after the references (none)

AFFILIATION = ("School of Computer Science and Engineering (SCOPE), Vellore Institute of "
               "Technology, Vellore, Tamil Nadu 632014, India")
COVER = {
    "title": ("Does Call-Graph Fan-Out Predict Failure Impact in Microservices? "
              "A Preregistered Fault-Injection Study"),
    # (name, affiliation number)
    "authors": [("Dakshrajsinh Tank", 1), ("Badrinath N", 1)],
    "affiliations": [AFFILIATION],
    "corresponding": "[TO FILL: corresponding author name and email]",
}

# Units shown in table headers. Time and percentage units come from the column-name suffix
# make_tables.py writes; count units name what is counted. Other columns are dimensionless
# (correlations, p-values, centralities) or categorical.
SUFFIX_UNITS = {"_ms": "ms", "_s": "s", "_pct": "%"}
COUNT_UNITS = {
    "n_services": "services", "n_services_actual": "services",
    "n_services_preregistered": "services", "ancestor_count": "services",
    "descendant_count": "services", "in_degree_raw": "callers",
    "out_degree_raw": "callees", "measured_ancestor_affected_kill": "services",
    "measured_ancestor_affected_latency": "services", "runs_invalidated": "runs",
    "tests": "tests", "tests_with_defined_p": "tests", "significant_after_bh_q05": "tests",
    "distinct_arrangements": "arrangements", "ancestor_count_levels_actual": "levels",
    "ancestor_count_levels_preregistered": "levels",
}
# Hazard ratios and their intervals are per one-unit increase of the predictor.
PER_UNIT = {c: "per unit of predictor" for c in ("cox_run_hr", "cox_run_ci_95",
                                                 "cox_service_hr", "cox_service_ci_95")}
TABLE_NOTE = ("Units are given in parentheses in the column headers; columns without a "
              "unit are dimensionless or categorical.")

# Footnotes for single tables, printed after TABLE_NOTE.
TABLE_FOOTNOTES = {
    "tbl_trec": ("Not estimable: the Cox fit did not converge or showed separation because "
                 "the predictor takes few distinct values; raw estimates are in the "
                 "supplementary file."),
}
# Supplementary data files copied into submission/: {published name: source under paper/}.
SUPPLEMENTARY = {
    "Supplementary_recovery_time_raw.csv": os.path.join("tables",
                                                        "tbl_trec_supplementary.csv"),
}

MARKER = re.compile(r"\s*%\[[^\]]*\]")
REF_MAP_ROW = re.compile(r"^>\s*\|\s*(Table|Figure)\s+(\d+)\s*\|\s*`([^`]+)`\s*\|")
MENTION = re.compile(r"\b(Tables?|Figures?)\s+(\d+)((?:\s*(?:,|and)\s*\d+)*)")
CITE = re.compile(r"\[(@[^\]\[]+)\]")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|<sub>.*?</sub>|<sup>.*?</sup>|\*[^*\s][^*]*\*)")


# ── drafts ────────────────────────────────────────────────────────────────────

def read(name):
    with open(os.path.join(DRAFT, name), encoding="utf-8") as fh:
        return fh.read()


def has_prose(name):
    return bool(name) and os.path.isfile(os.path.join(DRAFT, name))


def outline_heading(name, default):
    """The outline's own title, minus its ' — ... outline' suffix."""
    if not has_prose(name):
        return default
    first = read(name).splitlines()[0].lstrip("# ").strip()
    return re.split(r"\s+[—-]\s+", first)[0] or default


def draft_reference_map():
    """{('Table', n): file, ('Figure', n): file} from the results draft's drafting note."""
    m = {}
    for line in read("results.md").splitlines():
        r = REF_MAP_ROW.match(line)
        if r:
            m[(r.group(1), int(r.group(2)))] = r.group(3)
    return m


def blocks(text):
    """Split a draft into (kind, text) blocks, dropping drafting-note blockquotes."""
    out, para = [], []

    def flush():
        if para:
            out.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    lines = text.splitlines() + [""]
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i].lstrip("> ").rstrip())
                i += 1
            if "Drafting note" not in quote[0]:
                out.append(("quote", " ".join(q for q in quote if q)))
            continue
        h = re.match(r"^(#{1,3})\s+(.*)", line)
        if h:
            flush()
            out.append(("h%d" % len(h.group(1)), h.group(2).strip()))
        elif line.startswith("- "):
            flush()
            item = [line[2:].strip()]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip():
                item.append(lines[i].strip())
                i += 1
            out.append(("li", " ".join(item)))
            continue
        elif not line.strip():
            flush()
        else:
            para.append(line)
        i += 1
    return out


# ── references: name-year ─────────────────────────────────────────────────────

def load_refs(path=REFS):
    if not os.path.isfile(path):
        return {}
    with open(path, encoding="utf-8", newline="") as fh:
        return {r["key"]: r for r in csv.DictReader(fh)
                if r.get("status", "").strip() == "VERIFIED"}


def surnames(entry):
    """Authors are 'Surname, Initials; Surname, Initials; ...'."""
    return [a.split(",")[0].strip() for a in entry["authors"].split(";") if a.strip()]


def author_label(entry):
    s = surnames(entry)
    if len(s) == 1:
        return s[0]
    if len(s) == 2:
        return "%s & %s" % (s[0], s[1])
    if len(s) == 3:
        return "%s, %s & %s" % (s[0], s[1], s[2])
    return "%s et al." % s[0]


def year_suffixes(refs, keys):
    """Add a, b, ... to the year where two cited works share author label and year."""
    groups = {}
    for k in keys:
        groups.setdefault((author_label(refs[k]), refs[k]["year"]), []).append(k)
    out = {}
    for (_, year), ks in groups.items():
        ks = sorted(ks, key=lambda k: refs[k]["title"].lower())
        for i, k in enumerate(ks):
            out[k] = year + ("abcdefghijklmnopqrstuvwxyz"[i] if len(ks) > 1 else "")
    return out


def format_citation(keys, refs, years):
    entries = sorted(keys, key=lambda k: (years[k], author_label(refs[k]).lower()))
    return "(%s)" % "; ".join("%s, %s" % (author_label(refs[k]), years[k]) for k in entries)


def format_reference(entry, year):
    def short(a):
        sur, _, ini = a.partition(",")
        ini = re.sub(r"[.\s]", "", ini)
        return ("%s %s" % (sur.strip(), ini)).strip()
    authors = ", ".join(short(a) for a in entry["authors"].split(";") if a.strip())
    s = "%s. %s. %s. %s" % (authors, year, entry["title"].rstrip("."), entry["venue"])
    if entry.get("volume"):
        s += " %s" % entry["volume"]
    if entry.get("issue"):
        s += "(%s)" % entry["issue"]
    if entry.get("pages"):
        s += ":%s" % entry["pages"]
    s += "."
    if entry.get("doi"):
        s += " DOI %s." % entry["doi"]
    elif entry.get("url"):
        s += " Available at %s." % entry["url"]
    return s


def resolve_citations(text, refs, cited):
    """Replace [@a; @b] with a name-year citation when every key is VERIFIED."""
    def keys_of(m):
        parts = [p.strip() for p in m.group(1).split(";")]
        if not all(p.startswith("@") and len(p) > 1 for p in parts):
            return None
        return [p[1:] for p in parts]

    for m in CITE.finditer(text):
        ks = keys_of(m)
        if ks and all(k in refs for k in ks):
            for k in ks:
                if k not in cited:
                    cited.append(k)
    years = year_suffixes(refs, list(refs))

    def sub(m):
        ks = keys_of(m)
        if not ks or not all(k in refs for k in ks):
            return m.group(0)
        return format_citation(ks, refs, years)
    return CITE.sub(sub, text)


# ── styling ───────────────────────────────────────────────────────────────────

def style_document(doc, line_numbers=True):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12),
                       ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.5))
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1")
            ln.set(qn("w:restart"), "continuous")
            sec._sectPr.find(qn("w:pgMar")).addnext(ln)


def add_runs(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt(11)
        elif piece.startswith("<sub>"):
            par.add_run(piece[5:-6]).font.subscript = True
        elif piece.startswith("<sup>"):
            par.add_run(piece[5:-6]).font.superscript = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def placeholder(par, text="[TO WRITE]"):
    r = par.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0, 0)


def read_csv_rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.reader(fh))


def captions(index_file, key_col):
    rows = read_csv_rows(index_file)
    head = rows[0]
    k, s = head.index(key_col), head.index("shows")
    return {r[k]: r[s] for r in rows[1:]}


def header_with_unit(col):
    unit = COUNT_UNITS.get(col) or PER_UNIT.get(col)
    if unit is None:
        for suf, u in SUFFIX_UNITS.items():
            if col.endswith(suf):
                unit = u
                break
    return "%s (%s)" % (col, unit) if unit else col


# ── build ─────────────────────────────────────────────────────────────────────

class Builder:
    def __init__(self):
        self.doc = Document()
        style_document(self.doc)
        self.draft_map = draft_reference_map()
        self.tbl_caps = captions(os.path.join(HERE, "tables", "tables_index.csv"), "table")
        self.fig_caps = captions(os.path.join(HERE, "figures", "figures_index.csv"), "figure")
        self.number = {}          # file -> (kind, new number)
        self.counter = {"Table": 0, "Figure": 0}
        self.refs = load_refs()
        self.cited = []
        self.abstract_words = 0

    def resolve(self, kind, n):
        f = self.draft_map.get((kind, n))
        if f is None:
            return None
        if f not in self.number:
            self.counter[kind] += 1
            self.number[f] = (kind, self.counter[kind])
        return f

    def renumber(self, text):
        def sub(m):
            word, first, rest = m.group(1), int(m.group(2)), m.group(3) or ""
            kind = "Table" if word.startswith("Table") else "Figure"
            plural = word.endswith("s")
            nums = [first] + [int(x) for x in re.findall(r"\d+", rest)]
            new = []
            for n in nums:
                f = self.resolve(kind, n)
                new.append("%d [UNRESOLVED %s %d]" % (n, kind, n) if f is None
                           else str(self.number[f][1]))
            if kind == "Figure":
                before = text[:m.start()].rstrip()
                start = not before or before[-1] in ".!?"
                word = ("Figures" if plural else "Figure") if start else \
                       ("Figs." if plural else "Fig.")
            seps = re.findall(r"\s*(?:,|and)\s*", rest)
            out = word + " " + new[0]
            for sep, v in zip(seps, new[1:]):
                out += sep + v
            return out
        return MENTION.sub(sub, text)

    def paragraph(self, text, style=None):
        text = MARKER.sub("", text).strip()
        text = self.renumber(text)
        text = resolve_citations(text, self.refs, self.cited)
        add_runs(self.doc.add_paragraph(style=style), text)

    def abstract(self):
        self.doc.add_heading("Abstract", level=1)
        if has_prose(ABSTRACT_FILE):
            paras = [t for k, t in blocks(read(ABSTRACT_FILE)) if k == "p"]
        else:
            paras = ["**%s** [TO WRITE]" % h for h in ABSTRACT_PARTS]
        for t in paras:
            clean = MARKER.sub("", t).strip()
            self.abstract_words += len(re.sub(r"[*`]", "", clean).split())
            par = self.doc.add_paragraph()
            if "[TO WRITE]" in clean:
                head, _, rest = clean.partition("[TO WRITE]")
                add_runs(par, head)
                placeholder(par)
                add_runs(par, rest)
            else:
                add_runs(par, resolve_citations(self.renumber(clean), self.refs,
                                                self.cited))

    def to_write(self, heading):
        self.doc.add_heading(heading, level=1)
        placeholder(self.doc.add_paragraph())

    def build(self):
        self.doc.add_heading(COVER["title"], level=0)
        for heading, prose, outline in SECTIONS:
            if prose == ABSTRACT_FILE:
                self.abstract()
            elif has_prose(prose):
                for kind, text in blocks(read(prose)):
                    if kind.startswith("h"):
                        self.doc.add_heading(MARKER.sub("", text).strip(),
                                             level=int(kind[1]))
                    else:
                        self.paragraph(text, style={"quote": "Quote",
                                                    "li": "List Bullet"}.get(kind))
            else:
                self.to_write(outline_heading(outline, heading))

        self.doc.add_heading("References", level=1)
        if not self.refs:
            placeholder(self.doc.add_paragraph(),
                        "[TO WRITE: no reference in paper/references.csv has status "
                        "VERIFIED]")
        years = year_suffixes(self.refs, list(self.refs))
        for k in sorted(self.refs, key=lambda k: (surnames(self.refs[k])[0].lower(),
                                                  years[k], self.refs[k]["title"])):
            self.doc.add_paragraph(format_reference(self.refs[k], years[k]))
        for heading, prose, outline in TRAILING:
            self.to_write(outline_heading(outline, heading))

        cp = self.doc.core_properties
        cp.author, cp.title, cp.comments = "", "", "built by paper/build_docx.py"
        self.doc.save(OUT)

        # Figures and tables never mentioned in the text are numbered after the rest.
        for f in sorted(set(self.draft_map.values())):
            if f not in self.number:
                kind = "Table" if "/tables/" in f else "Figure"
                self.counter[kind] += 1
                self.number[f] = (kind, self.counter[kind])
        return self.write_submission()

    # ── submission files ──────────────────────────────────────────────────────
    def write_submission(self):
        os.makedirs(SUB, exist_ok=True)
        figs = sorted((n, f) for f, (k, n) in self.number.items() if k == "Figure")
        tbls = sorted((n, f) for f, (k, n) in self.number.items() if k == "Table")
        written = set()

        legends = Document()
        style_document(legends, line_numbers=False)
        legends.add_heading("Figure legends", level=1)
        from PIL import Image
        for n, f in figs:
            base = os.path.splitext(os.path.basename(f))[0]
            src_png = os.path.join(REPO, os.path.splitext(f)[0] + ".png")
            src_pdf = os.path.join(REPO, os.path.splitext(f)[0] + ".pdf")
            width = Image.open(src_png).size[0]
            if width < MIN_FIGURE_PX:
                raise SystemExit("%s is %d px wide; PeerJ needs at least %d"
                                 % (src_png, width, MIN_FIGURE_PX))
            for src, ext in ((src_png, "png"), (src_pdf, "pdf")):
                dst = "Figure %d.%s" % (n, ext)
                shutil.copyfile(src, os.path.join(SUB, dst))
                written.add(dst)
            p = legends.add_paragraph()
            p.add_run("Figure %d. " % n).bold = True
            p.add_run(self.fig_caps.get(base, base))
        legends.save(os.path.join(SUB, "figure_legends.docx"))
        written.add("figure_legends.docx")

        for n, f in tbls:
            base = os.path.splitext(os.path.basename(f))[0]
            rows = read_csv_rows(os.path.join(REPO, f))
            d = Document()
            style_document(d, line_numbers=False)
            if len(rows[0]) > 8:
                s = d.sections[0]
                s.orientation = WD_ORIENT.LANDSCAPE
                s.page_width, s.page_height = Inches(11), Inches(8.5)
            p = d.add_paragraph()
            p.add_run("Table %d. " % n).bold = True
            p.add_run(self.tbl_caps.get(base, base))
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = t.cell(i, j)
                    cell.text = ""
                    r = cell.paragraphs[0].add_run(header_with_unit(val) if i == 0 else val)
                    r.font.size = Pt(9)
                    r.bold = i == 0
                    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            d.add_paragraph(TABLE_NOTE)
            if base in TABLE_FOOTNOTES:
                d.add_paragraph(TABLE_FOOTNOTES[base])
            dst = "Table %d.docx" % n
            d.save(os.path.join(SUB, dst))
            written.add(dst)

        cover = Document()
        style_document(cover, line_numbers=False)
        cover.add_heading(COVER["title"], level=0)
        p = cover.add_paragraph()
        for i, (name, aff) in enumerate(COVER["authors"]):
            p.add_run(("" if i == 0 else ", ") + name)
            p.add_run(str(aff)).font.superscript = True
        for n, aff in enumerate(COVER["affiliations"], 1):
            p = cover.add_paragraph()
            p.add_run(str(n)).font.superscript = True
            p.add_run(" " + aff)
        cover.add_paragraph("Corresponding author: " + COVER["corresponding"])
        cover.save(os.path.join(SUB, "cover_page.docx"))
        written.add("cover_page.docx")

        if has_prose(COVER_LETTER):
            letter = Document()
            style_document(letter, line_numbers=False)
            letter.styles["Normal"].paragraph_format.line_spacing = 1.0
            for kind, text in blocks(read(COVER_LETTER)):
                if kind == "p":
                    text = (MARKER.sub("", text).strip().replace("{title}", COVER["title"])
                            .replace("{affiliation}", AFFILIATION))
                    add_runs(letter.add_paragraph(), text)
            letter.core_properties.author = ""
            letter.save(os.path.join(SUB, "cover_letter.docx"))
            written.add("cover_letter.docx")

        shutil.copyfile(OUT, os.path.join(SUB, "manuscript.docx"))
        written.add("manuscript.docx")
        for dst, src in SUPPLEMENTARY.items():
            shutil.copyfile(os.path.join(HERE, src), os.path.join(SUB, dst))
            written.add(dst)

        stale = sorted(set(os.listdir(SUB)) - written - LATER_FILES)
        return {"figures": figs, "tables": tbls, "stale": stale}


def main():
    b = Builder()
    info = b.build()
    print("wrote paper/manuscript.docx (no embedded figures or tables): %d verified "
          "references listed, %d cited in the rendered text"
          % (len(b.refs), len(b.cited)))
    uncited = sorted(set(b.refs) - set(b.cited))
    if uncited:
        print("  listed but not cited in the rendered text: %s" % ", ".join(uncited))
    print("abstract: %d words (limit %d)" % (b.abstract_words, ABSTRACT_MAX_WORDS))
    for kind in ("figures", "tables"):
        for n, f in info[kind]:
            print("  %s %d <- %s" % (kind[:-1].capitalize(), n, f))
    todo = [outline_heading(o, h) for h, p, o in SECTIONS + TRAILING if not has_prose(p)]
    print("  [TO WRITE]: %s" % "; ".join(todo))
    if info["stale"]:
        print("WARNING: paper/submission/ holds files this build did not write: %s"
              % ", ".join(info["stale"]))
    if b.abstract_words > ABSTRACT_MAX_WORDS:
        print("BUILD FAILED: abstract is %d words; the limit is %d"
              % (b.abstract_words, ABSTRACT_MAX_WORDS))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
