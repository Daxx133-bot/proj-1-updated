"""
build_review_pdf.py -- a review copy of the manuscript for internal feedback. It is not the
journal submission format and writes nothing to paper/submission/ or paper/manuscript.docx.

Outputs (paper/review/)
  Manuscript_review_copy.docx  plain Word styles, for comments and tracked changes
  Manuscript_review_copy.pdf   the same document exported by Microsoft Word
  review_checks.txt            the results of the checks below

Contents: a title page; a clickable table of contents; the abstract and sections 1 to 8;
Funding, Competing Interests, Data Availability, AI-Use Statement and References; and
Appendix A, the unrounded recovery-time estimates (paper/tables/tbl_trec_supplementary.csv,
transposed to fit the page). The deviations table is already in the body, so there is no
Appendix B.

Every body paragraph goes through paper/build_docx.py's own pipeline (markers stripped,
mentions renumbered, citations rendered), so the text is the submission text. Only placement
differs: each figure and table follows the paragraph that first cites it, with its number,
title and notes. Figures and tables are numbered as in the submission.

Layout: A4, 2.5 cm margins, Times New Roman 12 pt, 1.5 line spacing, justified body text,
line numbers in the margin, page numbers in the footer, landscape pages for tables wider than
seven columns. [TO FILL] placeholders are highlighted in yellow. Table cells carry zero-width
spaces after underscores so long identifiers can wrap; the checks ignore them.

Checks, written to review_checks.txt (the build fails if a hard check fails):
  text identity  the body paragraphs of the review docx equal the submission manuscript's, in
                 order, and each appears in order in the PDF's text layer (line numbers and
                 page numbers excluded)
  number audit   paper/check_docx_numbers.py's classification applied to the review docx;
                 table cells must equal their source CSV cells; 0 unmatched
  floats         every figure and table appears once, numbered in order, directly after the
                 paragraph that first cites it
  references     every VERIFIED reference is cited and listed
  layout         page count; text or table rules past the margins; a heading or table title
                 last on its page; a figure whose caption is on another page (reported only)

Requires Microsoft Word on Windows (via pywin32) for the table of contents, the table autofit
and the PDF export. Run after build_all.py: python paper/build_review_pdf.py
"""
from __future__ import annotations

import collections
import csv
import datetime
import os
import re
import sys
import unicodedata

import docx
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_docx as bd  # noqa: E402
import check_docx_numbers as cdn  # noqa: E402

HERE, REPO = bd.HERE, bd.REPO
OUTDIR = os.path.join(HERE, "review")
DOCX_OUT = os.path.join(OUTDIR, "Manuscript_review_copy.docx")
PDF_OUT = os.path.join(OUTDIR, "Manuscript_review_copy.pdf")
REPORT = os.path.join(OUTDIR, "review_checks.txt")

PAGE_W, PAGE_H, MARGIN = Cm(21.0), Cm(29.7), Cm(2.5)
TEXT_W, TEXT_H = Cm(16.0), Cm(24.7)
FIG_MAX_H = Cm(18.0)
LANDSCAPE_ABOVE_COLS = 7
REVIEW_NOTE = "Review copy for internal feedback, not the journal submission format"
PLACEHOLDER = re.compile(r"\[(?:TO FILL[^\]]*|Guide to confirm)\]")
ZWSP = "\u200b"
PAD = 6.5                    # pt: Word's default cell padding plus rounding
APPENDIX_A = {"heading": "Appendix A. Supplementary recovery-time estimates",
              "label": "Table A1",
              "file": "paper/tables/tbl_trec_supplementary.csv"}
CAPTION_LABEL = re.compile(r"^(Table|Figure) (A?\d+)\. ")
MENTION = re.compile(r"\b(Tables?|Figures?|Figs?\.)\s+(\d+)((?:\s*(?:,|and)\s*\d+)*)")


# ── docx helpers ──────────────────────────────────────────────────────────────

_plain_add_runs = bd.add_runs


def add_runs_highlighted(par, text):
    """build_docx.add_runs, with [TO FILL] placeholders highlighted in yellow."""
    pos = 0
    for m in PLACEHOLDER.finditer(text):
        _plain_add_runs(par, text[pos:m.start()])
        par.add_run(m.group(0)).font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    _plain_add_runs(par, text[pos:])


def add_field(par, instr, placeholder=""):
    def fld(kind):
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), kind)
        return e
    par.add_run()._r.append(fld("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " %s " % instr
    par.add_run()._r.append(it)
    par.add_run()._r.append(fld("separate"))
    par.add_run(placeholder)
    par.add_run()._r.append(fld("end"))


def page_setup(section, landscape=False, line_numbers=True):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = ((PAGE_H, PAGE_W) if landscape
                                               else (PAGE_W, PAGE_H))
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, MARGIN)
    sp = section._sectPr
    for e in sp.findall(qn("w:lnNumType")):
        sp.remove(e)
    if line_numbers:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sp.find(qn("w:pgMar")).addnext(ln)


def wrap_points(s):
    return s.replace("_", "_" + ZWSP)


def clean(s):
    return s.replace(ZWSP, "")


# ── build ─────────────────────────────────────────────────────────────────────

class ReviewBuilder(bd.Builder):
    def __init__(self):
        super().__init__()
        self.doc = Document()
        bd.style_document(self.doc, line_numbers=False)
        st = self.doc.styles
        st["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
            st[name].paragraph_format.keep_with_next = name != "Caption"
            # The template's theme-font attributes override the font name; drop them so
            # titles, headings and captions are Times New Roman like the body.
            rf = st[name].element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), bd.FONT)
        title_pPr = st["Title"].element.get_or_add_pPr()
        for e in title_pPr.findall(qn("w:pBdr")):
            title_pPr.remove(e)
        cap = st["Caption"]
        cap.font.name, cap.font.size = bd.FONT, Pt(10)
        cap.font.bold, cap.font.italic = False, False
        cap.font.color.rgb = RGBColor(0, 0, 0)
        cap.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), bd.FONT)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.line_spacing = 1.0
        self.landscape = False
        self.floats = []            # (label, kind, file, transposed) in document order
        self.unplaced = []

    # sections ---------------------------------------------------------------
    def start_section(self, landscape):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.different_first_page_header_footer = False
        page_setup(sec, landscape=landscape, line_numbers=True)
        self.landscape = landscape

    def ensure_orientation(self, landscape):
        if self.landscape != landscape:
            self.start_section(landscape)

    # floats -----------------------------------------------------------------
    def place_new(self, before):
        new = [f for f in self.number if f not in before]
        for f in new:
            kind, n = self.number[f]
            if kind == "Figure":
                self.add_figure("Figure %d" % n, f)
            else:
                self.add_table("Table %d" % n, f)
        if new:
            self.ensure_orientation(False)

    def add_figure(self, label, f):
        self.ensure_orientation(False)
        base = os.path.splitext(os.path.basename(f))[0]
        png = os.path.join(REPO, os.path.splitext(f)[0] + ".png")
        w, h = Image.open(png).size
        width = TEXT_W if TEXT_W * h / w <= FIG_MAX_H else int(FIG_MAX_H * w / h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(png, width=width)
        c = self.doc.add_paragraph(style="Caption")
        c.add_run(label + ". ").bold = True
        c.add_run(self.fig_caps.get(base, base))
        self.floats.append((label, "Figure", f, False))

    def add_table(self, label, f, transposed=False, caption=None):
        base = os.path.splitext(os.path.basename(f))[0]
        rows = bd.read_csv_rows(os.path.join(REPO, f))
        if transposed:
            rows = [list(col) for col in zip(*rows)]
        ncols = len(rows[0])
        self.ensure_orientation(ncols > LANDSCAPE_ABOVE_COLS)
        c = self.doc.add_paragraph(style="Caption")
        c.paragraph_format.keep_with_next = True
        c.add_run(label + ". ").bold = True
        c.add_run(caption or self.tbl_caps.get(base, base))
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.style = "Table Grid"
        t.autofit = True
        tblW = t._tbl.tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            t._tbl.tblPr.append(tblW)
        tblW.set(qn("w:type"), "pct")
        tblW.set(qn("w:w"), "5000")
        size = Pt(8) if ncols <= 10 else Pt(7)
        if transposed:
            # Unrounded values run to about 20 characters; narrower cell padding and a
            # slightly smaller font keep each value on one line.
            size = Pt(6.5)
            mar = OxmlElement("w:tblCellMar")
            for side in ("left", "right"):
                e = OxmlElement("w:" + side)
                e.set(qn("w:w"), "45")
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            t._tbl.tblPr.append(mar)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                head = (j == 0) if transposed else (i == 0)
                text = bd.header_with_unit(val) if head else val
                cell = t.cell(i, j)
                cell.text = ""
                par = cell.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(0)
                r = par.add_run(wrap_points(text))
                r.font.size = size
                r.bold = head or (transposed and i == 0)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trPr.append(hdr)
        self.doc.add_paragraph(bd.TABLE_NOTE, style="Caption")
        if base in bd.TABLE_FOOTNOTES:
            self.doc.add_paragraph(bd.TABLE_FOOTNOTES[base], style="Caption")
        self.floats.append((label, "Table", f, transposed))

    # text -------------------------------------------------------------------
    def paragraph(self, text, style=None):
        before = set(self.number)
        super().paragraph(text, style)
        self.place_new(before)

    def abstract(self):
        before = set(self.number)
        super().abstract()
        self.place_new(before)

    def title_page(self):
        doc = self.doc
        sec = doc.sections[0]
        page_setup(sec, landscape=False, line_numbers=False)
        sec.different_first_page_header_footer = True
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE", "1")

        doc.add_heading(bd.COVER["title"], level=0)
        p = doc.add_paragraph()
        for i, (name, aff) in enumerate(bd.COVER["authors"]):
            p.add_run(("" if i == 0 else ", ") + name)
            p.add_run(str(aff)).font.superscript = True
        for n, aff in enumerate(bd.COVER["affiliations"], 1):
            p = doc.add_paragraph()
            p.add_run(str(n)).font.superscript = True
            p.add_run(" " + aff)
        add_runs_highlighted(doc.add_paragraph(), "Corresponding author: "
                             + bd.COVER["corresponding"])
        today = datetime.date.today()
        doc.add_paragraph("%d %s" % (today.day, today.strftime("%B %Y")))
        r = doc.add_paragraph().add_run(REVIEW_NOTE)
        r.italic, r.font.size = True, Pt(10)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        r = doc.add_paragraph().add_run("Contents")
        r.bold, r.font.size = True, Pt(14)
        add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
                  "The table of contents is built when Word opens the document.")

    def build_review(self):
        bd.add_runs = add_runs_highlighted
        try:
            self.title_page()
            self.start_section(False)
            for heading, prose, outline in bd.SECTIONS:
                if prose == bd.ABSTRACT_FILE:
                    self.abstract()
                elif bd.has_prose(prose):
                    for kind, text in bd.blocks(bd.read(prose)):
                        if kind.startswith("h"):
                            self.doc.add_heading(bd.MARKER.sub("", text).strip(),
                                                 level=int(kind[1]))
                        else:
                            self.paragraph(text, style={"quote": "Quote",
                                                        "li": "List Bullet"}.get(kind))
                else:
                    self.to_write(bd.outline_heading(outline, heading))

            self.doc.add_heading("References", level=1)
            years = bd.year_suffixes(self.refs, list(self.refs))
            for k in sorted(self.refs, key=lambda k: (bd.surnames(self.refs[k])[0].lower(),
                                                      years[k], self.refs[k]["title"])):
                p = self.doc.add_paragraph(bd.format_reference(self.refs[k], years[k]))
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # The appendix heading goes on the table's own page, not alone before the break.
            ncols = len(bd.read_csv_rows(os.path.join(REPO, APPENDIX_A["file"])))
            self.ensure_orientation(ncols > LANDSCAPE_ABOVE_COLS)
            self.doc.add_heading(APPENDIX_A["heading"], level=1)
            self.add_table(APPENDIX_A["label"], APPENDIX_A["file"], transposed=True)
        finally:
            bd.add_runs = _plain_add_runs

        self.unplaced = sorted(f for f in set(self.draft_map.values()) if f not in self.number)
        cp = self.doc.core_properties
        cp.author, cp.title, cp.comments = "", "", "built by paper/build_review_pdf.py"
        os.makedirs(OUTDIR, exist_ok=True)
        self.doc.save(DOCX_OUT)


def word_finalize(docx_path, pdf_path):
    """Fit tables to the page, build the table of contents, save, and export the PDF."""
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(docx_path), False, False)
        for t in d.Tables:
            t.AutoFitBehavior(1)      # fit to contents
            t.AutoFitBehavior(2)      # then to the page width
        d.TablesOfContents(1).Update()
        d.Repaginate()
        d.TablesOfContents(1).UpdatePageNumbers()
        d.Save()
        d.ExportAsFixedFormat(OutputFileName=os.path.abspath(pdf_path), ExportFormat=17,
                              OpenAfterExport=False, OptimizeFor=0, CreateBookmarks=1,
                              DocStructureTags=True)
        pages = d.ComputeStatistics(2)
        d.Close(False)
    finally:
        word.Quit()
    return pages


# ── checks ────────────────────────────────────────────────────────────────────

def nfkc_compact(s):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", clean(s)))


def body_paragraphs(path, review):
    """(style, text) of the body: from Abstract to the end of References (review copy:
    captions, table notes, images and the appendix are left out)."""
    out, started = [], not review
    for p in docx.Document(path).paragraphs:
        t, st = clean(p.text).strip(), p.style.name
        if review:
            if not started:
                if st == "Heading 1" and t == "Abstract":
                    started = True
                else:
                    continue
            if st == "Heading 1" and t.startswith("Appendix"):
                break
            if st == "Caption":
                continue
        elif st == "Title":
            continue
        if t:
            out.append((st, t))
    return out


def check_identity(review_docx, pdf_path, L):
    sub = body_paragraphs(bd.OUT, review=False)
    rev = body_paragraphs(review_docx, review=True)
    diffs = []
    for i in range(max(len(sub), len(rev))):
        a = sub[i] if i < len(sub) else None
        b = rev[i] if i < len(rev) else None
        if a != b:
            diffs.append((i, a, b))
    L.append("Text identity, docx: %d submission paragraphs, %d review paragraphs, "
             "%d differences" % (len(sub), len(rev), len(diffs)))
    for i, a, b in diffs[:10]:
        L.append("  paragraph %d\n    submission: %r\n    review:     %r" % (i, a, b))

    import pdfplumber
    m = MARGIN.pt
    stream, body_started = [], False
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            words = page.extract_words(extra_attrs=["size"], use_text_flow=True)
            if not body_started:
                body_started = any(w["text"] == "Abstract" and w["size"] > 13 for w in words)
                if not body_started:
                    continue
            for w in words:
                if w["x1"] < m - 1 or w["top"] > page.height - m + 1:
                    continue        # line numbers in the left margin, page number in footer
                stream.append(w["text"])
    S = nfkc_compact("".join(stream))
    pos, missing = 0, []
    for st, t in sub:
        k = S.find(nfkc_compact(t), pos)
        if k < 0:
            missing.append(t)
        else:
            pos = k + len(nfkc_compact(t))
    L.append("Text identity, PDF text layer: %d of %d submission paragraphs found in order"
             % (len(sub) - len(missing), len(sub)))
    for t in missing[:10]:
        L.append("  not found: %s" % t[:160])
    return not diffs and not missing


def audit_review(review_docx, sources, L):
    nums, strings, by_id = cdn.load_numbers()
    sent_ids = cdn.draft_sentence_ids()
    counts = collections.Counter()
    flagged = []
    d = docx.Document(review_docx)
    pmap = {p._p: p for p in d.paragraphs}
    tmap = {t._tbl: t for t in d.tables}
    today = datetime.date.today()
    date_text = "%d %s" % (today.day, today.strftime("%B %Y"))
    section, label = "title page", None
    for el in d.element.body.iterchildren():
        if el in pmap:
            p = pmap[el]
            txt, st = clean(p.text).strip(), p.style.name
            if not txt:
                continue
            if st.lower().startswith("toc"):
                counts["table of contents entry"] += len(cdn.NUMBER.findall(txt))
                continue
            heading = st.startswith("Heading")
            if heading:
                section = txt
            m = CAPTION_LABEL.match(txt)
            if m:
                label = "%s %s" % m.groups()
            if section == "title page":
                txt = txt.replace(date_text, "")
                counts["date"] += 1 if date_text in clean(p.text) else 0
                if txt.endswith(bd.AFFILIATION) and txt[:-len(bd.AFFILIATION)].strip().isdigit():
                    counts["affiliation marker"] += 1
                    txt = bd.AFFILIATION
                txt = re.sub(r"(?<=[A-Za-z])\d$|(?<=[A-Za-z])\d(?=,)", "", txt)  # author marks
            if section == "References" and not heading:
                counts["reference entry"] += len(cdn.NUMBER.findall(cdn.norm(txt)))
                continue
            for sent in ([txt] if heading or st == "Title" else cdn.SENT.split(txt)):
                marked = [by_id[i] for i in sent_ids.get(cdn.key(sent), ()) if i in by_id]
                for tok, cat in cdn.classify_text(sent, heading or st == "Title", nums,
                                                  strings, 0, marked):
                    counts[cat] += 1
                    if cat in ("UNMATCHED", "WEAK MATCH"):
                        flagged.append((cat, section, tok, sent))
        elif el in tmap:
            src, transposed = sources.get(label, (None, False))
            srows = list(csv.reader(open(src, encoding="utf-8", newline=""))) if src else None
            for i, row in enumerate(tmap[el].rows):
                for j, c in enumerate(row.cells):
                    ct = clean(c.text)
                    toks = cdn.NUMBER.findall(cdn.norm(ct))
                    if not toks:
                        continue
                    a, b = (j, i) if transposed else (i, j)
                    head = (j == 0) if transposed else (i == 0)
                    same = (srows is not None and a < len(srows) and b < len(srows[a])
                            and srows[a][b] == ct)
                    cat = "table header" if head else "table cell" if same else "UNMATCHED"
                    counts[cat] += len(toks)
                    if cat == "UNMATCHED":
                        flagged.append((cat, label, ct, ""))
    L.append("Number audit (review docx): " + ", ".join(
        "%s %d" % (k, v) for k, v in sorted(counts.items())))
    for cat, sec, tok, sent in flagged:
        L.append("  %s [%s] %s | %s" % (cat, sec, tok, sent[:140]))
    return counts["UNMATCHED"]


def check_floats(review_docx, b, L):
    d = docx.Document(review_docx)
    seen, problems, first_cite = [], [], {}
    body_index = -1
    for p in d.paragraphs:
        txt, st = clean(p.text).strip(), p.style.name
        m = CAPTION_LABEL.match(txt) if st == "Caption" else None
        if m:
            kind, num = m.groups()
            seen.append((kind, num))
            if not num.startswith("A") and first_cite.get((kind, num)) != body_index:
                problems.append("%s %s is not directly after the paragraph that first "
                                "cites it" % (kind, num))
            continue
        if not txt or st == "Caption" or st.lower().startswith(("toc", "title", "heading")):
            continue
        body_index += 1
        for mm in MENTION.finditer(txt):
            k = "Table" if mm.group(1).startswith("Table") else "Figure"
            for n in [mm.group(2)] + re.findall(r"\d+", mm.group(3) or ""):
                first_cite.setdefault((k, n), body_index)
    for kind in ("Figure", "Table"):
        got = [n for k, n in seen if k == kind and not n.startswith("A")]
        want = [str(i) for i in range(1, len(got) + 1)]
        if got != want:
            problems.append("%s order %s, expected %s" % (kind, got, want))
    dup = [x for x, c in collections.Counter(seen).items() if c > 1]
    if dup:
        problems.append("appear more than once: %s" % dup)
    if b.unplaced:
        problems.append("never cited, not placed: %s" % b.unplaced)
    nfig = sum(1 for k, n in seen if k == "Figure")
    ntab = sum(1 for k, n in seen if k == "Table" and not n.startswith("A"))
    L.append("Floats: %d figures and %d tables in the body, each once, in citation order, "
             "directly after the paragraph that first cites it; appendix tables: %s; "
             "problems: %d" % (nfig, ntab, [k + " " + n for k, n in seen
                                            if n.startswith("A")], len(problems)))
    for x in problems:
        L.append("  " + x)
    return not problems


def check_references(review_docx, b, L):
    d = docx.Document(review_docx)
    listed, inside = 0, False
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            inside = p.text.strip() == "References"
            continue
        if inside and p.text.strip():
            listed += 1
    ok = listed == len(b.refs) and set(b.cited) == set(b.refs)
    L.append("References: %d VERIFIED, %d listed, %d cited" % (len(b.refs), listed,
                                                               len(b.cited)))
    return ok


def check_layout(review_docx, pdf_path, L):
    import pdfplumber
    headings = {nfkc_compact(p.text) for p in docx.Document(review_docx).paragraphs
                if p.style.name.startswith("Heading") and p.text.strip()}
    m = MARGIN.pt
    issues = []
    with pdfplumber.open(pdf_path) as pdf:
        n = len(pdf.pages)
        for no, page in enumerate(pdf.pages, 1):
            W, H = page.width, page.height
            words = [w for w in page.extract_words(use_text_flow=True)
                     if w["top"] < H - m + 1]
            over = [w["text"] for w in words if w["x1"] > W - m + 2]
            # Word draws table borders one cell padding (0.08 in, 5.76 pt) outside the text
            # margins; only rules beyond that count.
            rules = [e for e in page.edges if e["x1"] > W - m + PAD or e["x0"] < m - PAD]
            if over or rules:
                issues.append("page %d: content past the right or left margin (%s)"
                              % (no, ", ".join(over[:5]) or "table rules"))
            lines = collections.OrderedDict()
            for w in words:
                if w["x1"] < m - 1:
                    continue
                lines.setdefault(round(w["top"]), []).append(w["text"])
            if not lines:
                continue
            last = "".join(lines[sorted(lines)[-1]])
            if nfkc_compact(last) in headings:
                issues.append("page %d: heading '%s' is the last line" % (no, last))
            if CAPTION_LABEL.match(" ".join(lines[sorted(lines)[-1]])):
                issues.append("page %d: a figure or table title is the last line" % no)
            for img in page.images:
                below = [t for top, t in lines.items() if top > img["bottom"]]
                if not any(" ".join(t).startswith("Figure") for t in below):
                    issues.append("page %d: a figure without its caption below it" % no)
    L.append("Layout: %d pages; problems: %d" % (n, len(issues)))
    for x in issues:
        L.append("  " + x)
    return n, issues


def main():
    b = ReviewBuilder()
    b.build_review()
    pages = word_finalize(DOCX_OUT, PDF_OUT)
    sources = {label: (os.path.join(REPO, f), tr) for label, kind, f, tr in b.floats
               if kind == "Table"}
    L = ["review_checks.txt -- written by paper/build_review_pdf.py",
         "Review copy: paper/review/Manuscript_review_copy.pdf (%d pages in Word) and .docx"
         % pages, ""]
    ok_text = check_identity(DOCX_OUT, PDF_OUT, L)
    unmatched = audit_review(DOCX_OUT, sources, L)
    ok_floats = check_floats(DOCX_OUT, b, L)
    ok_refs = check_references(DOCX_OUT, b, L)
    n, issues = check_layout(DOCX_OUT, PDF_OUT, L)
    todo = [(p.style.name, m.group(0)) for p in docx.Document(DOCX_OUT).paragraphs
            for m in PLACEHOLDER.finditer(p.text)]
    L.append("Placeholders (highlighted): %s" % ("; ".join(t for _, t in todo) or "none"))
    with open(REPORT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    print("\n".join(L))
    failed = [name for name, ok in (("text identity", ok_text),
                                    ("number audit", unmatched == 0),
                                    ("floats", ok_floats), ("references", ok_refs)) if not ok]
    if failed:
        print("REVIEW BUILD FAILED: %s" % ", ".join(failed))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
